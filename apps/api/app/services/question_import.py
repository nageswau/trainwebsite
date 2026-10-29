"""nehn-025: the trainer's standard question template and the row rules of its import.

`template_xlsx` builds the download: a header-only `Questions` sheet first (the sheet `read_tabular` reads), then an `Instructions` sheet
with the column reference and worked examples, so nothing has to be deleted before uploading. `parse_rows` turns `read_tabular` rows into
`AssessmentQuestionIn` objects -- the shape the single and bulk question routes take, so an imported MCQ grades exactly like a typed one --
plus one error per bad row and the non-blocking warnings. The route decides what to do with them (all or nothing)."""

import io
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from app.schemas import AssessmentQuestionIn
from app.services.keyword_scoring import KEYWORD_THRESHOLD_PERCENT, MAX_KEYWORDS

TYPES = ("mcq_single", "mcq_multiple", "text", "file")
OPTION_COLUMNS = tuple(f"option_{n}" for n in range(1, 7))
COLUMNS = ("type", "question", *OPTION_COLUMNS, "correct", "marks", "keywords", "explanation")
REQUIRED = ("type", "question")
MAX_ROWS = 500  # the bulk question route's limit
MAX_MARKS = 10000
NUMBERS = re.compile(r"\d+(\s*,\s*\d+)*")

REFERENCE = (
    ("type", "Yes", "mcq_single (one right answer), mcq_multiple (one or more), text (written answer) or file (the student uploads a file)", "mcq_single"),
    ("question", "Yes", "The question text, 2 to 10,000 characters", "What does HTML stand for?"),
    ("option_1 … option_6", "For MCQs", "The answer choices; a multiple-choice question needs at least 2. Leave them empty for text and file questions.", "HyperText Markup Language"),
    ("correct", "For MCQs", "The number of the right option, e.g. 2; for mcq_multiple list every right option, e.g. 1,3. Empty for text and file questions.", "1,3"),
    ("marks", "No", "Whole number from 1 to 10,000; empty means 1", "2"),
    (
        "keywords",
        "No",
        f"Text questions only: up to {MAX_KEYWORDS} keywords separated by commas; put synonyms together with | (loop|iteration). An answer "
        f"with {KEYWORD_THRESHOLD_PERCENT}% of them gets a suggested full score (fewer: in proportion) that you review before publishing.",
        "for loop, loop|iteration",
    ),
    ("explanation", "No", "Reserved for an answer explanation; not used yet", ""),
)
EXAMPLES = (
    ("mcq_single", "What does HTML stand for?", "HyperText Markup Language", "High Tech Modern Language", "", "", "", "", "1", "1"),
    ("mcq_multiple", "Which of these are programming languages?", "Python", "HTML", "Java", "CSS", "", "", "1,3", "2"),
    ("text", "Explain the difference between a list and a tuple in Python.", "", "", "", "", "", "", "", "5"),
    ("file", "Upload a screenshot of your working program.", "", "", "", "", "", "", "", "10"),
)
NOTES = (
    "Fill in the Questions sheet (the first sheet) -- one question per row, below the header. Do not rename or reorder the header row.",
    f"Up to {MAX_ROWS} questions per file. Empty rows are skipped. A .csv file with the same header row is accepted too.",
    "Images, formulas and formatting inside a question are not supported -- type the question as plain text (Unicode and symbols such as √ π ∑ are fine).",
    "The upload is checked first: every row with a problem is listed by its row number and nothing is imported until the file has no errors.",
    "Imported questions are added after the assessment's existing questions. A Draft assessment stays hidden from students until you publish it.",
)


def template_xlsx() -> bytes:
    book = openpyxl.Workbook()
    questions = book.active
    questions.title = "Questions"
    questions.append(list(COLUMNS))
    for cell in questions[1]:
        cell.font = Font(bold=True)
    questions.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHIJKL", (14, 50, 18, 18, 18, 18, 18, 18, 10, 8, 20, 30), strict=True):
        questions.column_dimensions[letter].width = width
    kinds = DataValidation(type="list", formula1=f'"{",".join(TYPES)}"', allow_blank=True)
    kinds.add(f"A2:A{MAX_ROWS + 1}")
    questions.add_data_validation(kinds)

    guide = book.create_sheet("Instructions")
    guide.append(["How to fill in the Questions sheet"])
    guide["A1"].font = Font(bold=True, size=13)
    for note in NOTES:
        guide.append([note])
    guide.append([])
    guide.append(["Column", "Required", "What to enter", "Example"])
    for cell in guide[guide.max_row]:
        cell.font = Font(bold=True)
    for line in REFERENCE:
        guide.append(list(line))
    guide.append([])
    guide.append(["Worked examples (copy the rows you need into the Questions sheet)"])
    guide.cell(guide.max_row, 1).font = Font(bold=True)
    guide.append(list(COLUMNS[:10]))
    for cell in guide[guide.max_row]:
        cell.font = Font(bold=True)
    for example in EXAMPLES:
        guide.append(list(example))
    for letter, width in zip("ABCD", (22, 50, 60, 30), strict=True):
        guide.column_dimensions[letter].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _span(numbers: list[int]) -> str:
    """[1, 2, 3, 4] -> "1-4"; [1, 3] -> "1, 3"."""
    if len(numbers) > 1 and numbers == list(range(numbers[0], numbers[-1] + 1)):
        return f"{numbers[0]}-{numbers[-1]}"
    return ", ".join(map(str, numbers))


def _question(values: dict[str, str]) -> AssessmentQuestionIn | str:
    """The question a row describes, or the first problem with it."""
    kind = values.get("type", "").lower()
    prompt = values.get("question", "")
    if not kind:
        return "Type is required"
    if kind not in TYPES:
        return f"Type must be one of {', '.join(TYPES)}"
    if not prompt:
        return "Question is required"
    if not 2 <= len(prompt) <= 10000:
        return "Question must be 2 to 10,000 characters"
    filled = {n: values[column] for n, column in enumerate(OPTION_COLUMNS, start=1) if values.get(column)}
    correct = values.get("correct", "")
    if kind.startswith("mcq"):
        if len(filled) < 2:
            return "A multiple-choice question needs at least 2 options"
        seen: dict[str, int] = {}
        for n, text in filled.items():
            if text.casefold() in seen:
                return f"Options must be different (option {n} repeats option {seen[text.casefold()]})"
            seen[text.casefold()] = n
        if not correct:
            return "Correct is required for a multiple-choice question (e.g. 2 or 1,3)"
        if not NUMBERS.fullmatch(correct):
            return "Correct must be option numbers separated by commas (e.g. 1,3)"
        chosen = sorted({int(n) for n in correct.split(",")})
        for n in chosen:
            if n not in filled:
                return f"Correct answer {n} is not a filled option (options {_span(list(filled))} are filled)"
        if kind == "mcq_single" and len(chosen) > 1:
            return "A mcq_single question has exactly one correct answer"
    elif filled or correct:
        return f"A {kind} question takes no options or correct answer"
    else:
        chosen = []
    marks = values.get("marks", "") or "1"
    if not marks.isdigit() or not 1 <= int(marks) <= MAX_MARKS:
        return "Marks must be a whole number from 1 to 10,000"
    keywords = values.get("keywords", "").split(",")  # nehn-026: the schema trims them and owns the keyword rules and their messages
    try:
        return AssessmentQuestionIn(question_type=kind, prompt=prompt, options=list(filled.values()), correct_answers=[filled[n] for n in chosen], max_score=int(marks), keywords=keywords)
    except ValidationError as error:
        return error.errors()[0]["msg"]


def parse_rows(rows: list[tuple[int, dict[str, str]]]) -> tuple[list[tuple[int, AssessmentQuestionIn]], list[dict], list[dict]]:
    """(row, question) for every good row; one {row, message} error per bad row; the warnings (repeated questions, unsaved columns)."""
    questions, errors, warnings = [], [], []
    first_seen: dict[str, int] = {}
    unsaved = 0
    for number, values in rows:
        result = _question(values)
        if isinstance(result, str):
            errors.append({"row": number, "message": result})
        else:
            questions.append((number, result))
        key = " ".join(values.get("question", "").split()).casefold()
        if key and key in first_seen:
            warnings.append({"row": number, "message": f"Repeats the question in row {first_seen[key]}"})
        first_seen.setdefault(key, number)
        unsaved += bool(values.get("explanation"))
    if unsaved:
        fill = "row fills" if unsaved == 1 else "rows fill"
        warnings.append({"row": None, "message": f"The explanation column is not saved yet; {unsaved} {fill} it in."})
    return questions, errors, warnings
