"""rec-012 (DEC-SCOPE-148, spec §3): rule-based resume extraction (R7: in-house, no AI provider). Pure: bytes in, text out; text and
the Skills Master's terms in, suggestions out -- no database and no I/O, so the caller decides scope and what is stored.

Untrusted files: both readers are pure Python and run nothing embedded. A PDF is read for its first MAX_PAGES pages only; a DOCX is
refused before parsing when its ZIP holds too many parts or too many uncompressed bytes (a zip bomb), and python-docx parses the XML
with entity resolution off. The caller runs this in a thread under a time limit (EX8)."""

import io
import re
import zipfile
from collections.abc import Hashable, Iterable

from app.services.candidates import PDF

MAX_PAGES = 30
MAX_CHARS = 100_000
MAX_DOCX_PARTS = 2_000
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_ITEMS = 10
ENCRYPTED = "This PDF is password-protected. Upload a copy without a password."
UNREADABLE = "Could not read the text of this resume. Check that the file opens, or add the details by hand."
TOO_LARGE = "This DOCX is too large to read. Save a smaller copy and upload it again."


class ExtractError(Exception):
    """A file this module will not read; the message is a sentence for the recruiter."""


# --- reading ---------------------------------------------------------------------------------------------------------------------
def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader  # imported here: the module stays importable for the pure matcher

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")  # many "protected" PDFs only restrict printing and open with an empty password
        except Exception:  # an unsupported cipher is as unreadable to us as a real password
            opened = 0
        if not opened:
            raise ExtractError(ENCRYPTED)
    return "\n".join(page.extract_text() or "" for page in reader.pages[:MAX_PAGES])


def _docx_text(data: bytes) -> str:
    from docx import Document

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        parts = archive.infolist()
        if len(parts) > MAX_DOCX_PARTS or sum(p.file_size for p in parts) > MAX_DOCX_UNCOMPRESSED:
            raise ExtractError(TOO_LARGE)
    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # resumes often lay skills out in a table
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if not cells or cells[-1] != cell.text:  # a merged cell repeats in python-docx
                    cells.append(cell.text)
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_text(data: bytes, content_type: str) -> tuple[str, bool]:
    """The resume's text (stripped, at most MAX_CHARS) and whether it was cut. A scanned PDF has no text layer and answers ''."""
    try:
        text = _pdf_text(data) if content_type == PDF else _docx_text(data)
    except ExtractError:
        raise
    except Exception:  # every parser failure on an untrusted file is the same sentence, never a 500
        raise ExtractError(UNREADABLE) from None
    text = text.replace("\x00", "").strip()
    return text[:MAX_CHARS], len(text) > MAX_CHARS


# --- skills (EX3, EX4) -----------------------------------------------------------------------------------------------------------
# A token starts with a letter or digit (or "." / "#" for .NET) and may hold + # . inside (C++, C#, Node.js); a trailing "." is the
# sentence's, not the token's. The text and every term go through this one tokeniser, so "CI/CD" matches "CI/CD".
_TOKEN = re.compile(r"[.#]?[A-Za-z0-9][A-Za-z0-9+#.]*")
# Ordinary English words that are also skill names: they match only in the master's own casing ("Go", not "go"). Single letters too.
STOP_WORDS = frozenset({"go", "spring", "rest", "swift", "rust", "express", "shell", "less", "access", "excel"})


def _tokens(text: str) -> list[tuple[str, int, int]]:
    out = []
    for m in _TOKEN.finditer(text):
        word = m.group().rstrip(".")
        if word:
            out.append((word, m.start(), m.start() + len(word)))
    return out


def _index(terms: Iterable[tuple[str, Hashable]]) -> dict[tuple[str, ...], list[tuple[Hashable, tuple[str, ...] | None]]]:
    """Lower-cased token tuple -> [(skill id, the exact tokens a case-sensitive term needs, else None)]."""
    index: dict = {}
    for term, skill_id in terms:
        words = tuple(w for w, _, _ in _tokens(term))
        if not words:
            continue
        exact = words if len(words) == 1 and (len(words[0]) == 1 or words[0].lower() in STOP_WORDS) else None
        index.setdefault(tuple(w.lower() for w in words), []).append((skill_id, exact))
    return index


def _hit(index: dict, words: list[str]) -> Hashable | None:
    """The skill these words name: exactly, or with a plural "s" on the last word ("REST APIs" -> REST API)."""
    lowered = [w.lower() for w in words]
    candidates = [(lowered, words)]
    if len(lowered[-1]) > 2 and lowered[-1].endswith("s"):
        candidates.append((lowered[:-1] + [lowered[-1][:-1]], words[:-1] + [words[-1][:-1]]))
    for key, original in candidates:
        for skill_id, exact in index.get(tuple(key), ()):
            if exact is None or tuple(original) == exact:
                return skill_id
    return None


def match_skills(text: str, terms: Iterable[tuple[str, Hashable]]) -> list[dict]:
    """Leftmost-longest matching: at each token the longest term wins and its tokens are consumed, so "Spring Boot" is not also
    "Spring". One entry per skill (a name and its alias count once), in order of first appearance, with the text as written."""
    index = _index(terms)
    if not index:
        return []
    longest = max(len(k) for k in index)
    tokens = _tokens(text)
    found: dict = {}
    i = 0
    while i < len(tokens):
        for n in range(min(longest, len(tokens) - i), 0, -1):
            skill_id = _hit(index, [w for w, _, _ in tokens[i : i + n]])
            if skill_id is not None:
                found.setdefault(skill_id, text[tokens[i][1] : tokens[i + n - 1][2]])
                i += n
                break
        else:
            i += 1
    return [{"skill_id": skill_id, "matched": matched} for skill_id, matched in found.items()]


# --- other suggestions (EX5) -----------------------------------------------------------------------------------------------------
_EXPERIENCE = re.compile(
    r"(?<![\d.])(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years?|yrs?)(?:\s*(?:and\s*)?(\d{1,2})\s*(?:months?|mos?))?\b"
    r"|(?<![\d.])(\d{1,2})\s*(?:months?|mos?)\b",
    re.IGNORECASE,
)
_NEAR = 25  # characters either side in which "experience" makes a figure the total


def experience_months(text: str) -> int | None:
    """The total experience in months: the first figure next to the word "experience", else the first figure at all."""
    first = None
    for m in _EXPERIENCE.finditer(text):
        if m.group(3):
            months = int(m.group(3))
        else:
            months = round(float(m.group(1)) * 12) + int(m.group(2) or 0)
        if not 0 < months <= 600:
            continue
        window = text[max(0, m.start() - _NEAR) : m.end() + _NEAR].lower()
        if "experien" in window:
            return months
        first = first if first is not None else months
    return first


_NOT_INSIDE = r"(?<![\w@.])"  # not the tail of a word, an email or a domain ("web.com" is not B.Com)
_DEGREES = [  # (label, rank, pattern); the highest rank named wins, the first named on a tie
    ("Ph.D.", 5, re.compile(_NOT_INSIDE + r"(?:ph\.?\s?d\b\.?|doctorate\b)", re.I)),
    ("M.Tech", 4, re.compile(_NOT_INSIDE + r"(?:m\.?\s?tech\b|master of technology\b)", re.I)),
    ("M.E.", 4, re.compile(_NOT_INSIDE + r"(?:M\.E\.|[Mm]aster of [Ee]ngineering\b)")),
    ("MCA", 4, re.compile(_NOT_INSIDE + r"(?:mca\b|master of computer applications?\b)", re.I)),
    ("MBA", 4, re.compile(_NOT_INSIDE + r"(?:mba\b|master of business administration\b)", re.I)),
    ("M.Sc", 4, re.compile(_NOT_INSIDE + r"(?:m\.?\s?sc\b|master of science\b)", re.I)),
    ("M.Com", 4, re.compile(_NOT_INSIDE + r"(?:m\.\s?com\b|mcom\b|master of commerce\b)", re.I)),
    ("M.A.", 4, re.compile(_NOT_INSIDE + r"master of arts\b", re.I)),
    ("B.Tech", 3, re.compile(_NOT_INSIDE + r"(?:b\.?\s?tech\b|bachelor of technology\b)", re.I)),
    ("B.E.", 3, re.compile(_NOT_INSIDE + r"(?:B\.E\.|[Bb]achelor of [Ee]ngineering\b)")),
    ("BCA", 3, re.compile(_NOT_INSIDE + r"(?:bca\b|bachelor of computer applications?\b)", re.I)),
    ("BBA", 3, re.compile(_NOT_INSIDE + r"(?:bba\b|bachelor of business administration\b)", re.I)),
    ("B.Sc", 3, re.compile(_NOT_INSIDE + r"(?:b\.?\s?sc\b|bachelor of science\b)", re.I)),
    ("B.Com", 3, re.compile(_NOT_INSIDE + r"(?:b\.\s?com\b|bcom\b|bachelor of commerce\b)", re.I)),
    ("B.A.", 3, re.compile(_NOT_INSIDE + r"bachelor of arts\b", re.I)),
    ("Diploma", 2, re.compile(r"\bdiploma\b", re.I)),
]


def qualification(text: str) -> str | None:
    best = None  # (rank, -position, label)
    for label, rank, pattern in _DEGREES:
        m = pattern.search(text)
        if m and (best is None or (rank, -m.start()) > best[:2]):
            best = (rank, -m.start(), label)
    return best[2] if best else None


_LOCATION_LINE = re.compile(r"^[ \t]*(?:current location|location|city|based in)[ \t]*[:\-–][ \t]*(.+?)[ \t]*$", re.I | re.M)
CITIES = (
    "Hyderabad", "Secunderabad", "Bengaluru", "Bangalore", "Chennai", "Mumbai", "Navi Mumbai", "Pune", "New Delhi", "Delhi", "Noida",
    "Gurugram", "Gurgaon", "Kolkata", "Ahmedabad", "Kochi", "Coimbatore", "Visakhapatnam", "Vijayawada", "Guntur", "Tirupati", "Warangal",
    "Jaipur", "Indore", "Chandigarh", "Lucknow", "Bhubaneswar", "Nagpur", "Mysuru", "Mysore", "Thiruvananthapuram", "Trivandrum",
    "Mangaluru", "Madurai", "Vadodara", "Surat", "Bhopal", "Patna", "Ranchi", "Raipur", "Dehradun", "Goa",
)  # fmt: skip
_CITY = re.compile(r"\b(" + "|".join(sorted((re.escape(c) for c in CITIES), key=len, reverse=True)) + r")\b", re.I)
_CITY_NAMES = {c.lower(): c for c in CITIES}


def location(text: str) -> str | None:
    """A "Location:" line as written, else the first known city named."""
    if m := _LOCATION_LINE.search(text):
        return m.group(1)[:120]
    if m := _CITY.search(text):
        return _CITY_NAMES[m.group(1).lower()]
    return None


_TITLE = re.compile(
    r"\b(?:(?:Senior|Sr\.?|Junior|Jr\.?|Lead|Principal|Associate|Staff)\s+)?"
    r"(?:(?:Software|Java|Python|Full[ -]?Stack|Front[ -]?End|Back[ -]?End|Web|Data|DevOps|QA|Test|Cloud|Mobile|Android|iOS|\.NET|React|"
    r"Node(?:\.js)?|Salesforce|SAP|Database|Network|Systems?|Business|ML|AI|Machine Learning|UI|UX)\s+)+"
    r"(?:Developer|Engineer|Architect|Analyst|Consultant|Tester|Programmer|Administrator|Scientist|Intern|Designer)\b",
    re.I,
)
_CERT_ACRONYM = re.compile(r"\b(OCJP|OCA|OCP|SCJP|CCNA|CCNP|PMP|CSM|PSM|ITIL|CISSP|CEH|RHCE|RHCSA|CKA|CKAD)\b")
_CERT_LABEL = re.compile(r"^(?:certifications?|certificates?)\s*[:\-–]\s*(.*)$", re.I)
_BULLET = re.compile(r"^[\s\-•*·▪●◦]+")


def _unique(values: Iterable[str]) -> list[str]:
    seen, out = set(), []
    for value in values:
        value = value.strip(" \t,;:-–")
        if value and value.lower() not in seen:
            seen.add(value.lower())
            out.append(value)
    return out[:MAX_ITEMS]


def job_titles(text: str) -> list[str]:
    return _unique(" ".join(m.group().split()) for m in _TITLE.finditer(text))


def certifications(text: str) -> list[str]:
    """Lines that mention a certification (a "Certifications:" line is split into its items), then the known acronyms anywhere."""
    found = []
    for raw in text.splitlines():
        line = _BULLET.sub("", raw).strip()
        if not line or len(line) > 150 or "certifi" not in line.lower():
            continue
        if m := _CERT_LABEL.match(line):
            found.extend(re.split(r"[,;|]", m.group(1)))
        else:
            found.append(line)
    found.extend(m.group() for m in _CERT_ACRONYM.finditer(text))
    return _unique(found)


INDUSTRIES = {
    "Banking & Finance": ("banking", "bank", "fintech", "financial services", "capital markets"),
    "Insurance": ("insurance",),
    "Healthcare": ("healthcare", "hospital", "pharma", "pharmaceutical", "clinical"),
    "E-commerce & Retail": ("e-commerce", "ecommerce", "retail"),
    "Telecom": ("telecom", "telecommunications"),
    "EdTech": ("edtech", "e-learning"),  # not "education": every resume has an Education section
    "Manufacturing": ("manufacturing", "automotive", "automobile"),
    "Logistics": ("logistics", "supply chain"),
    "Travel & Hospitality": ("travel", "hospitality", "airline"),
}
_INDUSTRY = {name: re.compile(r"\b(?:" + "|".join(re.escape(w) for w in words) + r")\b", re.I) for name, words in INDUSTRIES.items()}


def industries(text: str) -> list[str]:
    """The industries named, in order of first mention."""
    hits = [(m.start(), name) for name, pattern in _INDUSTRY.items() if (m := pattern.search(text))]
    return [name for _, name in sorted(hits)][:5]


def suggest(text: str, terms: Iterable[tuple[str, Hashable]]) -> dict:
    """Every suggestion for one resume's text. `terms` are (skill name or alias, skill id) for the active skills."""
    return {
        "skills": match_skills(text, terms),
        "qualification": qualification(text),
        "experience_months": experience_months(text),
        "location": location(text),
        "job_titles": job_titles(text),
        "certifications": certifications(text),
        "industries": industries(text),
    }
