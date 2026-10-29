"""nehn-025 -- trainer question upload from a standard Excel template.

`GET /workflows/it/trainer/assessments/question-template.xlsx` and `POST .../assessments/{id}/questions/import` (a dry run by default,
`?commit=true` to write). AC1 the filled template imports 20 mixed questions with 0 errors into a Draft; AC2 a bad `correct` names its row
and nothing is written; AC3 imported MCQs auto-grade like manual ones; AC4 students never see the Draft."""

import datetime
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func, select

from app.core.security import hash_password
from app.models import Assessment, AssessmentAttempt, AssessmentQuestion, AuditLog, Batch, Enrollment, Program, User

PASSWORD = "Sup3r-Secret-Pass!"
BASE = "/api/v1/workflows/it/trainer/assessments"
TEMPLATE = f"{BASE}/question-template.xlsx"
HEADER = ["type", "question", "option_1", "option_2", "option_3", "option_4", "option_5", "option_6", "correct", "marks", "keywords", "explanation"]


async def _user(db, role: str, division: str = "it") -> User:
    user = User(email=f"n025-{role}-{uuid.uuid4().hex[:8]}@example.local", password_hash=hash_password(PASSWORD), full_name=f"N025 {role}", role=role, division=division, active=True)
    db.add(user)
    await db.commit()
    return user


async def _batch(db, trainer: User) -> Batch:
    program = Program(
        slug=f"n025-{uuid.uuid4().hex[:8]}",
        category="Software Development",
        title=f"N025 Program {uuid.uuid4().hex[:6]}",
        summary="Test",
        duration="8 weeks",
        eligibility="None",
        fees=10000,
        certification="Test cert",
        curriculum=["Module 1"],
        placement_assistance="Yes",
        trainer_name="Test Trainer",
        active=True,
    )
    db.add(program)
    await db.flush()
    batch = Batch(
        program_id=program.id,
        trainer_id=trainer.id,
        name=f"N025-{uuid.uuid4().hex[:6]}",
        start_date=datetime.date.today(),
        end_date=datetime.date.today() + datetime.timedelta(days=90),
        schedule="Mon-Fri 7pm",
        capacity=20,
        enrollment_open=True,
        status="active",
    )
    db.add(batch)
    await db.commit()
    return batch


async def _assessment(db, batch: Batch, status: str = "draft") -> Assessment:
    item = Assessment(batch_id=batch.id, title=f"N025 Quiz {uuid.uuid4().hex[:6]}", scheduled_at=datetime.datetime.now(datetime.UTC), status=status)
    db.add(item)
    await db.commit()
    return item


async def _login(client, user: User) -> None:
    response = await client.post("/api/v1/auth/login", json={"email": user.email, "password": PASSWORD, "division": user.division})
    assert response.status_code == 200, response.text


def _xlsx(rows: list[list], header: list[str] = HEADER) -> bytes:
    book = openpyxl.Workbook()
    book.active.append(header)
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _row(kind: str, question: str, options: list[str] = (), correct="", marks=None, keywords="", explanation="") -> list:
    padded = list(options) + [None] * (6 - len(options))
    return [kind, question, *padded, correct, marks, keywords, explanation]


def _twenty() -> list[list]:
    """AC1: 20 mixed questions -- single, multiple (numeric and comma answers), text and file, with Unicode and maths symbols."""
    rows = []
    for n in range(1, 9):
        rows.append(_row("mcq_single", f"Q{n}: What is {n} + {n}?", [str(n), str(2 * n), str(3 * n), str(4 * n)], 2, 2))
    for n in range(9, 15):
        rows.append(_row("mcq_multiple", f"Q{n}: Pick the even numbers (n={n})", ["1", "2", "3", "4"], "2, 4"))
    for n in range(15, 19):
        rows.append(_row("text", f"Q{n}: Explain √{n} and π in தமிழ்", marks=5))
    rows.append(_row("file", "Q19: Upload your diagram", marks=10))
    rows.append(_row("MCQ_SINGLE", "Q20: ∑ of 1..3 is?", ["5", "6"], "2"))
    return rows


async def _import(client, assessment_id, raw: bytes, *, commit: bool = False, filename: str = "questions.xlsx"):
    return await client.post(f"{BASE}/{assessment_id}/questions/import", params={"commit": str(commit).lower()}, files={"file": (filename, raw, "application/octet-stream")})


async def _question_count(db, assessment_id) -> int:
    return await db.scalar(select(func.count()).select_from(AssessmentQuestion).where(AssessmentQuestion.assessment_id == assessment_id))


# --- template ---------------------------------------------------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_template_is_an_xlsx_with_the_columns_and_an_instructions_sheet(db_session, client):
    trainer = await _user(db_session, "trainer")
    await _login(client, trainer)
    response = await client.get(TEMPLATE)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="question-template.xlsx"' in response.headers["content-disposition"]
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "Questions" and "Instructions" in book.sheetnames
    assert [c.value for c in next(book["Questions"].iter_rows(max_row=1))] == HEADER
    assert book["Questions"].max_row == 1  # header only: nothing to delete before uploading
    instructions = " ".join(str(c.value) for row in book["Instructions"].iter_rows() for c in row if c.value)
    assert "mcq_multiple" in instructions and "1,3" in instructions and "not supported" in instructions.lower()


@pytest.mark.asyncio
@pytest.mark.parametrize("role,division", [("it_student", "it"), ("counselor", "overseas")])
async def test_template_is_refused_to_other_roles(db_session, client, role, division):
    await _login(client, await _user(db_session, role, division))
    assert (await client.get(TEMPLATE)).status_code == 403


@pytest.mark.asyncio
async def test_template_requires_sign_in(client):
    assert (await client.get(TEMPLATE)).status_code == 401


# --- AC1: dry run, then commit --------------------------------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_ac1_twenty_mixed_questions_preview_with_no_errors_and_write_nothing(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    response = await _import(client, assessment.id, _xlsx(_twenty() + [[None] * 12, [None] * 12]))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["total_rows"] == 20 and body["errors"] == [] and body["created"] == 0 and len(body["questions"]) == 20
    first = body["questions"][0]
    assert first == {"row": 2, "question_type": "mcq_single", "prompt": "Q1: What is 1 + 1?", "options": ["1", "2", "3", "4"], "correct_answers": ["2"], "max_score": 2}
    assert body["questions"][8]["correct_answers"] == ["2", "4"]
    assert body["questions"][14] == {"row": 16, "question_type": "text", "prompt": "Q15: Explain √15 and π in தமிழ்", "options": [], "correct_answers": [], "max_score": 5}
    assert body["questions"][19]["question_type"] == "mcq_single"  # the type is not case-sensitive
    assert await _question_count(db_session, assessment.id) == 0


@pytest.mark.asyncio
async def test_ac1_commit_creates_every_question_in_row_order_after_the_existing_ones_and_stays_a_draft(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    db_session.add(AssessmentQuestion(assessment_id=assessment.id, question_type="text", prompt="Typed earlier", position=4))
    await db_session.commit()
    await _login(client, trainer)
    response = await _import(client, assessment.id, _xlsx(_twenty()), commit=True)
    assert response.status_code == 201, response.text
    assert response.json()["created"] == 20 and response.json()["errors"] == []
    rows = (
        await db_session.execute(select(AssessmentQuestion.prompt, AssessmentQuestion.position).where(AssessmentQuestion.assessment_id == assessment.id).order_by(AssessmentQuestion.position))
    ).all()
    assert [r.prompt for r in rows][:2] == ["Typed earlier", "Q1: What is 1 + 1?"]
    assert [r.position for r in rows] == list(range(4, 25))
    assert await db_session.scalar(select(Assessment.status).where(Assessment.id == assessment.id)) == "draft"
    audit = (await db_session.scalars(select(AuditLog).where(AuditLog.entity_id == str(assessment.id), AuditLog.action == "assessment.questions_import"))).all()
    assert len(audit) == 1 and audit[0].metadata_json == {"count": 20, "format": "xlsx"}


@pytest.mark.asyncio
async def test_a_csv_with_the_same_columns_is_accepted(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    csv = ",".join(HEADER) + "\r\nmcq_single,Capital of France?,Paris,Rome,,,,,1,,,\r\n"
    response = await _import(client, assessment.id, csv.encode(), commit=True, filename="q.csv")
    assert response.status_code == 201, response.text
    assert response.json()["created"] == 1


# --- AC2: row errors, all or nothing ------------------------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_ac2_a_correct_answer_outside_the_options_names_its_row_and_commits_nothing(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    rows = _twenty()[:3] + [_row("mcq_single", "Bad one", ["a", "b", "c", "d"], 7)]
    preview = await _import(client, assessment.id, _xlsx(rows))
    assert preview.status_code == 200
    assert preview.json()["errors"] == [{"row": 5, "message": "Correct answer 7 is not a filled option (options 1-4 are filled)"}]
    commit = await _import(client, assessment.id, _xlsx(rows), commit=True)
    assert commit.status_code == 422
    assert commit.json()["detail"] == "1 row has errors -- fix them and upload the file again. Nothing was imported."
    assert await _question_count(db_session, assessment.id) == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "row,message",
    [
        (_row("essay", "Q?"), "Type must be one of mcq_single, mcq_multiple, text, file"),
        (_row("", "Q?"), "Type is required"),
        (_row("text", ""), "Question is required"),
        (_row("text", "Q"), "Question must be 2 to 10,000 characters"),
        (_row("mcq_single", "Q?", ["only"], 1), "A multiple-choice question needs at least 2 options"),
        (_row("mcq_single", "Q?", ["a", "b"]), "Correct is required for a multiple-choice question (e.g. 2 or 1,3)"),
        (_row("mcq_single", "Q?", ["a", "b"], "1,2"), "A mcq_single question has exactly one correct answer"),
        (_row("mcq_multiple", "Q?", ["a", "b"], "1;x"), "Correct must be option numbers separated by commas (e.g. 1,3)"),
        (_row("mcq_single", "Q?", ["a", "A ", "b"], 1), "Options must be different (option 2 repeats option 1)"),
        (_row("mcq_single", "Q?", ["a", None, "c"], 2), "Correct answer 2 is not a filled option (options 1, 3 are filled)"),
        (_row("text", "Q?", ["a"]), "A text question takes no options or correct answer"),
        (_row("file", "Q?", correct=1), "A file question takes no options or correct answer"),
        (_row("text", "Q?", marks=0), "Marks must be a whole number from 1 to 10,000"),
        (_row("text", "Q?", marks="2.5"), "Marks must be a whole number from 1 to 10,000"),
    ],
)
async def test_row_rules(db_session, client, row, message):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    response = await _import(client, assessment.id, _xlsx([row]))
    assert response.status_code == 200, response.text
    assert response.json()["errors"] == [{"row": 2, "message": message}]
    assert response.json()["questions"] == []


@pytest.mark.asyncio
async def test_options_may_skip_a_column_and_are_stored_by_text(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    response = await _import(client, assessment.id, _xlsx([_row("mcq_multiple", "Q?", ["a", None, "c", "d"], "4,1,1")]))
    assert response.json()["errors"] == []
    assert response.json()["questions"][0]["options"] == ["a", "c", "d"] and response.json()["questions"][0]["correct_answers"] == ["a", "d"]


@pytest.mark.asyncio
async def test_duplicates_and_unsaved_columns_are_warnings_not_errors(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    rows = [_row("text", "Same question", keywords="alpha, beta"), _row("text", " same QUESTION ", explanation="Because")]
    body = (await _import(client, assessment.id, _xlsx(rows))).json()
    assert body["errors"] == []
    assert body["warnings"] == [
        {"row": 3, "message": "Repeats the question in row 2"},
        {"row": None, "message": "The keywords and explanation columns are not saved yet; 2 rows fill them in."},
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "header,status,detail",
    [
        (["question", "answer"], 422, "Unknown column: answer; Missing required column: type"),
        (["type", "question", "type"], 422, "Duplicate column: type"),
    ],
)
async def test_a_wrong_header_is_a_clear_file_error(db_session, client, header, status, detail):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    response = await _import(client, assessment.id, _xlsx([], header=header))
    assert response.status_code == status and response.json()["detail"] == detail


@pytest.mark.asyncio
async def test_file_level_refusals(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer))
    await _login(client, trainer)
    empty = await _import(client, assessment.id, _xlsx([[None] * 12]))
    assert empty.status_code == 422 and empty.json()["detail"] == "The file has no questions"
    too_many = await _import(client, assessment.id, _xlsx([_row("text", f"Question {n}") for n in range(501)]))
    assert too_many.status_code == 422 and too_many.json()["detail"] == "The file has more than 500 rows"
    wrong_type = await _import(client, assessment.id, b"%PDF-1.4", filename="q.pdf")
    assert wrong_type.status_code == 415
    not_xlsx = await _import(client, assessment.id, b"not a zip", filename="q.xlsx")
    assert not_xlsx.status_code == 422 and not_xlsx.json()["detail"] == "The file is not a valid Excel (.xlsx) file"


# --- scope ------------------------------------------------------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_a_trainer_cannot_import_into_another_trainers_assessment(db_session, client):
    owner = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, owner))
    await _login(client, await _user(db_session, "trainer"))
    assert (await _import(client, assessment.id, _xlsx(_twenty()))).status_code == 403
    assert (await _import(client, assessment.id, _xlsx(_twenty()), commit=True)).status_code == 403
    assert await _question_count(db_session, assessment.id) == 0


@pytest.mark.asyncio
async def test_it_admin_may_import_into_any_assessment(db_session, client):
    assessment = await _assessment(db_session, await _batch(db_session, await _user(db_session, "trainer")))
    await _login(client, await _user(db_session, "it_admin"))
    assert (await _import(client, assessment.id, _xlsx(_twenty()), commit=True)).status_code == 201


@pytest.mark.asyncio
@pytest.mark.parametrize("role,division", [("it_student", "it"), ("counselor", "overseas")])
async def test_other_roles_cannot_import(db_session, client, role, division):
    assessment = await _assessment(db_session, await _batch(db_session, await _user(db_session, "trainer")))
    await _login(client, await _user(db_session, role, division))
    assert (await _import(client, assessment.id, _xlsx(_twenty()), commit=True)).status_code == 403


@pytest.mark.asyncio
async def test_unknown_assessment_is_404_and_sign_in_is_required(db_session, client):
    assert (await _import(client, uuid.uuid4(), _xlsx(_twenty()))).status_code == 401
    await _login(client, await _user(db_session, "trainer"))
    assert (await _import(client, uuid.uuid4(), _xlsx(_twenty()))).status_code == 404


@pytest.mark.asyncio
async def test_an_assessment_a_student_has_started_is_locked(db_session, client):
    trainer = await _user(db_session, "trainer")
    assessment = await _assessment(db_session, await _batch(db_session, trainer), status="scheduled")
    db_session.add(AssessmentAttempt(assessment_id=assessment.id, student_id=(await _user(db_session, "it_student")).id, status="in_progress"))
    await db_session.commit()
    await _login(client, trainer)
    assert (await _import(client, assessment.id, _xlsx(_twenty()))).status_code == 409
    assert (await _import(client, assessment.id, _xlsx(_twenty()), commit=True)).status_code == 409


# --- AC3 + AC4: students ------------------------------------------------------------------------------------------------------
@pytest.mark.asyncio
async def test_ac3_ac4_imported_mcqs_grade_like_manual_ones_and_the_draft_stays_hidden(db_session, client):
    trainer = await _user(db_session, "trainer")
    batch = await _batch(db_session, trainer)
    student = await _user(db_session, "it_student")
    db_session.add(Enrollment(student_id=student.id, batch_id=batch.id, enrollment_code=f"N025-{uuid.uuid4().hex[:6]}", status="active"))
    await db_session.commit()
    imported = await _assessment(db_session, batch)
    manual = await _assessment(db_session, batch)
    rows = [_row("mcq_single", "Pick b", ["a", "b"], 2, 3), _row("mcq_multiple", "Pick a and c", ["a", "b", "c"], "1,3", 2)]
    await _login(client, trainer)
    assert (await _import(client, imported.id, _xlsx(rows), commit=True)).status_code == 201
    for payload in (
        {"question_type": "mcq_single", "prompt": "Pick b", "options": ["a", "b"], "correct_answers": ["b"], "max_score": 3},
        {"question_type": "mcq_multiple", "prompt": "Pick a and c", "options": ["a", "b", "c"], "correct_answers": ["a", "c"], "max_score": 2, "position": 2},
    ):
        assert (await client.post(f"{BASE}/{manual.id}/questions", json=payload)).status_code == 201

    # AC4: while both are Drafts the student neither sees nor starts them
    await _login(client, student)
    titles = {row["title"] for row in (await client.get("/api/v1/portal/it/student/examinations")).json()["rows"]}
    assert imported.title not in titles
    assert (await client.post(f"/api/v1/workflows/it/assessments/{imported.id}/attempts")).status_code == 404

    await _login(client, trainer)
    for item in (imported, manual):
        assert (await client.patch(f"{BASE}/{item.id}", json={"status": "scheduled"})).status_code == 200

    await _login(client, student)
    assert imported.title in {row["title"] for row in (await client.get("/api/v1/portal/it/student/examinations")).json()["rows"]}
    results = []
    for item in (imported, manual):
        attempt = (await client.post(f"/api/v1/workflows/it/assessments/{item.id}/attempts")).json()
        by_prompt = {q["prompt"]: q["id"] for q in attempt["questions"]}
        answers = [{"question_id": by_prompt["Pick b"], "value": "b"}, {"question_id": by_prompt["Pick a and c"], "value": ["c", "a"]}]
        submitted = await client.post(f"/api/v1/workflows/it/assessment-attempts/{attempt['id']}/submit", json={"answers": answers})
        assert submitted.status_code == 200, submitted.text
        results.append({k: submitted.json()[k] for k in ("status", "score", "percentage", "grade")})
    assert results[0] == results[1] == {"status": "graded", "score": 5.0, "percentage": 100.0, "grade": results[0]["grade"]}
