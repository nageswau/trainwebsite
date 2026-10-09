"""rec-012 (DEC-SCOPE-148, spec §3): the pure resume extractor -- bytes in, text out; text and terms in, suggestions out. No database."""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.services import resume_extract as rx
from app.services.candidates import DOCX, PDF

TERMS = [
    ("Java", "java"), ("J2EE", "java"), ("Spring", "spring"), ("Spring Boot", "boot"), ("Hibernate", "hibernate"), ("REST API", "rest"),
    ("MySQL", "mysql"), ("Go", "go"), ("C", "c"), ("C++", "cpp"), ("C#", "csharp"), ("Node.js", "node"), ("CI/CD", "cicd"), ("Git", "git"),
]  # fmt: skip
SOURCE = "Developed enterprise applications using Java, Spring Boot, Hibernate and REST APIs with MySQL."


def _ids(result: dict) -> list:
    return [s["skill_id"] for s in result["skills"]]


def _pdf(*lines: str, pages: int = 1) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for page in range(pages):
        for i, line in enumerate(lines):
            pdf.drawString(40, 800 - 14 * i, line.replace("{page}", str(page + 1)))
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- skills (AC1, EX4) -----------------------------------------------------------------------------------------------------------
def test_source_sentence_yields_exactly_the_five_skills():
    """AC1: Java, Spring Boot, Hibernate, REST API ("REST APIs") and MySQL -- and not Spring, which Spring Boot already covers."""
    result = rx.suggest(SOURCE, TERMS)
    assert _ids(result) == ["java", "boot", "hibernate", "rest", "mysql"]
    assert [s["matched"] for s in result["skills"]] == ["Java", "Spring Boot", "Hibernate", "REST APIs", "MySQL"]


def test_alias_and_name_of_one_skill_suggest_it_once():
    assert _ids(rx.suggest("J2EE developer. Strong Java.", TERMS)) == ["java"]


def test_symbols_and_slashes_match_whole_terms():
    result = rx.suggest("Skills: C++, C#; Node.js. Pipelines (CI/CD) with Git.", TERMS)
    assert _ids(result) == ["cpp", "csharp", "node", "cicd", "git"]


def test_terms_never_match_inside_words():
    assert _ids(rx.suggest("Javascript, Mysqlx, Gitlab, Hibernated", TERMS)) == []


def test_stop_words_and_single_letters_need_the_masters_casing():
    """Edge case: "go" in prose is not Go; "c" is not C. The master's own casing still matches."""
    assert _ids(rx.suggest("Ready to go the extra mile; grade c; rest assured; in spring.", TERMS)) == []
    assert _ids(rx.suggest("Languages: Go, C. Frameworks: Spring.", TERMS)) == ["go", "c", "spring"]


def test_matching_ignores_case_for_ordinary_terms():
    assert _ids(rx.suggest("JAVA and mysql", TERMS)) == ["java", "mysql"]


def test_no_terms_or_no_text_yields_nothing():
    assert rx.suggest("", TERMS)["skills"] == []
    assert rx.suggest(SOURCE, [])["skills"] == []


# --- other suggestions (EX5) -----------------------------------------------------------------------------------------------------
@pytest.mark.parametrize(
    ("text", "months"),
    [
        ("3 years of experience in Java", 36),
        ("Total Experience: 4.5 yrs", 54),
        ("Experience - 2 years 6 months", 30),
        ("8 months experience as intern", 8),
        ("Worked 5+ years in banking", 60),
        ("No numbers here", None),
        ("Built 120 years of nothing", None),  # more than 50 years is not a career
    ],
)
def test_experience(text, months):
    assert rx.suggest(text, [])["experience_months"] == months


def test_experience_prefers_the_figure_next_to_the_word_experience():
    assert rx.suggest("Led a 2 years project. Overall 6 years of experience.", [])["experience_months"] == 72


@pytest.mark.parametrize(
    ("text", "qualification"),
    [
        ("B.Tech in Computer Science, 2021", "B.Tech"),
        ("Education: MCA from Osmania University; BCA 2018", "MCA"),
        ("Bachelor of Engineering (Mechanical)", "B.E."),
        ("M.Sc Computer Science", "M.Sc"),
        ("Ph.D. in Physics; M.Tech", "Ph.D."),
        ("Diploma in Electronics", "Diploma"),
        ("Self taught", None),
    ],
)
def test_qualification_is_the_highest_degree_named(text, qualification):
    assert rx.suggest(text, [])["qualification"] == qualification


def test_location_from_a_label_then_a_known_city():
    assert rx.suggest("Name: Ravi\nLocation: Hyderabad, Telangana\nSkills", [])["location"] == "Hyderabad, Telangana"
    assert rx.suggest("Ravi Kumar | ravi@example.com | Bengaluru", [])["location"] == "Bengaluru"
    assert rx.suggest("Nothing about places", [])["location"] is None


def test_job_titles_certifications_and_industry():
    text = (
        "Senior Java Developer at Acme Bank (2021-2024)\nSoftware Engineer at ShopKart\n"
        "AWS Certified Solutions Architect - Associate\nCertifications: OCJP, CCNA\n"
        "Domain: banking and e-commerce; insurance clients."
    )
    result = rx.suggest(text, [])
    assert result["job_titles"] == ["Senior Java Developer", "Software Engineer"]
    assert "AWS Certified Solutions Architect - Associate" in result["certifications"]
    assert {"OCJP", "CCNA"} <= set(result["certifications"])
    assert result["industries"] == ["Banking & Finance", "E-commerce & Retail", "Insurance"]


def test_a_certification_shares_a_line_with_other_parts():
    """QA-01: a header line "Title | Certification" yields only the certification part."""
    result = rx.suggest("Senior Java Developer | AWS Certified Developer - Associate", [])
    assert result["certifications"] == ["AWS Certified Developer - Associate"]


def test_lists_are_capped():
    text = "\n".join(f"Software Engineer {i}\n" for i in range(30)) + "\n".join(f"Certified Thing {i}" for i in range(30))
    result = rx.suggest(text, [])
    assert len(result["certifications"]) <= rx.MAX_ITEMS
    assert len(result["job_titles"]) <= rx.MAX_ITEMS


# --- reading files (AC3, EX8, EX9) -----------------------------------------------------------------------------------------------
def test_pdf_text_feeds_the_matcher():
    text, truncated = rx.read_text(_pdf(SOURCE), PDF)
    assert not truncated
    assert _ids(rx.suggest(text, TERMS)) == ["java", "boot", "hibernate", "rest", "mysql"]


def test_docx_paragraphs_and_tables_and_three_years():
    """Positive scenario: a DOCX resume yields the skills plus "3 years"."""
    data = _docx("Ravi Kumar", "3 years of experience building services.", table=[["Skills", "Java, Spring Boot, MySQL"]])
    text, _ = rx.read_text(data, DOCX)
    result = rx.suggest(text, TERMS)
    assert _ids(result) == ["java", "boot", "mysql"]
    assert result["experience_months"] == 36


def test_scanned_or_empty_pdf_has_no_text():
    """AC3: a PDF page with only a drawing has no text layer."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.rect(40, 40, 200, 200, fill=1)
    pdf.showPage()
    pdf.save()
    assert rx.read_text(buffer.getvalue(), PDF) == ("", False)


def test_encrypted_pdf_is_a_readable_error():
    """Negative scenario: an encrypted PDF is an ExtractError with a sentence, never an unhandled exception."""
    writer = PdfWriter()
    writer.append(io.BytesIO(_pdf(SOURCE)))
    writer.encrypt(user_password="secret", owner_password="owner")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(rx.ExtractError, match="password-protected"):
        rx.read_text(buffer.getvalue(), PDF)


def test_pdf_with_an_empty_user_password_is_read():
    writer = PdfWriter()
    writer.append(io.BytesIO(_pdf(SOURCE)))
    writer.encrypt(user_password="", owner_password="owner")
    buffer = io.BytesIO()
    writer.write(buffer)
    assert "Hibernate" in rx.read_text(buffer.getvalue(), PDF)[0]


@pytest.mark.parametrize("content_type", [PDF, DOCX])
def test_corrupt_files_are_a_readable_error(content_type):
    data = b"%PDF-1.7 not really" if content_type == PDF else b"PK\x03\x04garbage"
    with pytest.raises(rx.ExtractError, match="Could not read"):
        rx.read_text(data, content_type)


def test_only_the_first_pages_are_read():
    text, _ = rx.read_text(_pdf("Page {page} marker", pages=rx.MAX_PAGES + 2), PDF)
    assert f"Page {rx.MAX_PAGES} marker" in text
    assert f"Page {rx.MAX_PAGES + 1} marker" not in text


def test_very_long_text_is_truncated_at_the_cap(monkeypatch):
    """Edge case: a very long resume is cut at the cap and says so."""
    monkeypatch.setattr(rx, "MAX_CHARS", 50)
    text, truncated = rx.read_text(_docx("word " * 100), DOCX)
    assert truncated
    assert len(text) == 50


def test_docx_zip_bomb_guard(monkeypatch):
    monkeypatch.setattr(rx, "MAX_DOCX_UNCOMPRESSED", 1000)
    data = _docx("x" * 5000)
    with pytest.raises(rx.ExtractError, match="too large"):
        rx.read_text(data, DOCX)


def test_docx_with_too_many_parts(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", "<x/>")
        for i in range(5):
            archive.writestr(f"junk/{i}", "")
    monkeypatch.setattr(rx, "MAX_DOCX_PARTS", 3)
    with pytest.raises(rx.ExtractError, match="too large"):
        rx.read_text(buffer.getvalue(), DOCX)
