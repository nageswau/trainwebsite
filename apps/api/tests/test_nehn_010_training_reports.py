"""nehn-010 (DEC-SCOPE-305) -- the training reports pack in the nehn-009 registry: (a) student-wise course report, (b) students per
trainer, (c) course completion, (d) completions in a period, (e) attendance, (f) trainer report, (g) course-wise report with revenue
(`financial`). One fixture course with two batches and two trainers; every figure is checked against hand-counted values, and the
attendance % against the certificate eligibility endpoint (AC2)."""

import datetime
import io
import uuid
from datetime import UTC
from datetime import datetime as dt

import openpyxl
import pytest

from app.core import rbac
from app.models import Assessment, AssessmentAttempt, Assignment, Attendance, Certificate, Payment, Submission
from app.services.reports import registry
from tests.bdm001_helpers import login, make_user
from tests.test_nehn_009_reports import _course, _enrol

pytestmark = pytest.mark.asyncio
API = "/api/v1/reports"
PACK = ("student-courses", "students-by-trainer", "course-completion", "completions", "attendance", "trainers", "courses")
DAY = datetime.date(2020, 1, 6)
AT = dt(2020, 1, 20, 10, tzinfo=UTC)


async def _fixture(db):
    """Course P: batch B1 (trainer T1, ends 2020-02-01) and B2 (trainer T2, ends 2020-03-01).
    s1 in B1: progress 100, 6 classes (2 present, late, excused, 2 absent → 66.67 %), assignments 1 graded / 2 published (+1 draft,
    1 still `submitted`), assessments 1 passed / 2 counted (+1 draft), certificate issued, fee 1000 paid.
    s2 in B1: progress 40, no attendance, fee 1000 pending. s3 in B1: dropped at progress 100. s1's student again in B2: progress 50,
    fee 500 `succeeded` (an `enrollment` legacy reference)."""
    program, b1, t1 = await _course(db)
    _, b2, t2 = await _course(db, end=datetime.date(2020, 3, 1))
    b2.program_id = program.id
    s1, student = await _enrol(db, b1, progress=100, name="Aaron N010")
    s2, _ = await _enrol(db, b1, progress=40, name="Bella N010")
    s3, _ = await _enrol(db, b1, progress=100, status="dropped", name="Chen N010")
    s4 = (await _enrol(db, b2, progress=50, name="Dummy"))[0]
    s4.student_id = student.id
    for day, status in enumerate(("present", "present", "late", "excused", "absent", "absent")):
        db.add(Attendance(student_id=student.id, batch_id=b1.id, session_date=DAY + datetime.timedelta(days=day), status=status))
    graded, pending, draft = (Assignment(batch_id=b1.id, title=t, description="-", due_date=AT, published=p) for t, p in (("G", True), ("P", True), ("D", False)))
    db.add_all([graded, pending, draft])
    passed, failed, unscheduled = (Assessment(batch_id=b1.id, title=t, scheduled_at=AT, pass_percent=p, status=s) for t, p, s in (("X1", 50, "published"), ("X2", 60, "closed"), ("X3", 50, "draft")))
    db.add_all([passed, failed, unscheduled])
    await db.flush()
    db.add_all([
        Submission(assignment_id=graded.id, student_id=student.id, status="graded", score=90),
        Submission(assignment_id=pending.id, student_id=student.id, status="submitted"),
        AssessmentAttempt(assessment_id=passed.id, student_id=student.id, attempt_no=1, status="graded", percentage=40),
        AssessmentAttempt(assessment_id=passed.id, student_id=student.id, attempt_no=2, status="graded", percentage=70),
        AssessmentAttempt(assessment_id=failed.id, student_id=student.id, attempt_no=1, status="submitted", percentage=80),
        Certificate(student_id=student.id, program_id=program.id, enrollment_id=s1.id, certificate_no=f"N010-{uuid.uuid4().hex[:8]}",
                    verification_code=f"N010-{uuid.uuid4().hex[:8]}", issued_on=datetime.date(2020, 2, 5), status="issued"),
        Payment(user_id=student.id, division="it", reference_type="enrollment_fee", reference_id=s1.id, amount=1000, status="paid"),
        Payment(user_id=s2.student_id, division="it", reference_type="enrollment_fee", reference_id=s2.id, amount=1000, status="pending"),
        Payment(user_id=student.id, division="it", reference_type="enrollment", reference_id=s4.id, amount=500, status="succeeded"),
    ])
    await db.commit()
    return {"program": program, "b1": b1, "b2": b2, "t1": t1, "t2": t2, "s1": s1, "s2": s2, "s3": s3, "s4": s4}


async def _as(client, db, role):
    user = await make_user(db, role, "global" if role == "super_admin" else "it")
    await login(client, user)
    return user


async def _read(client, report, **params) -> dict:
    response = await client.get(f"{API}/{report}", params={"page_size": 200, **params})
    assert response.status_code == 200, response.text
    return response.json()


def _by(body, key) -> dict:
    return {row[key]: row for row in body["items"]}


# --- registration (AC1) ----------------------------------------------------------------------------------------------------------


def test_the_pack_is_registered_under_the_training_resource():
    for report_id in PACK:
        report = registry.REPORTS[report_id]
        assert report.resource == "reports.training" and "recruiter_id" not in report.filters
    assert registry.REPORTS["courses"].classes == {"revenue": "financial"}
    assert all(not registry.REPORTS[r].classes for r in PACK if r != "courses")
    assert "field.financial.view" not in rbac.DEFAULT_GRANTS["it_admin"] + rbac.DEFAULT_GRANTS["trainer"]


async def test_the_catalogue_lists_the_pack_for_it_admin(client, db_session):
    await _as(client, db_session, "it_admin")
    ids = [item["id"] for item in (await client.get(API)).json()["items"]]
    assert set(PACK) <= set(ids) and ids[0] == "enrollments"


# --- the figures ------------------------------------------------------------------------------------------------------------------


async def test_a_student_wise_course_report_carries_the_eligibility_criteria(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    body = await _read(client, "student-courses", course_id=str(f["program"].id))
    assert [c["key"] for c in body["columns"]] == [
        "student_code", "student", "course", "batch", "trainer", "enrolled_on", "status", "progress", "attendance_pct",
        "assignments_graded", "assignments_total", "assessments_passed", "assessments_total", "certificate",
    ]
    assert body["total"] == 4  # the student in two batches of one course is two rows
    rows = {(r["student"], r["batch"]): r for r in body["items"]}
    s1 = rows[("Aaron N010", f["b1"].name)]
    assert (s1["progress"], s1["attendance_pct"], s1["assignments_graded"], s1["assignments_total"], s1["assessments_passed"], s1["assessments_total"]) == (100, 66.67, 1, 2, 1, 2)
    assert s1["certificate"] == "issued" and s1["trainer"] == f["t1"].full_name
    s2 = rows[("Bella N010", f["b1"].name)]
    assert s2["attendance_pct"] is None and s2["certificate"] == "Not issued"  # 0/0 is "—", not 0 %
    assert rows[("Aaron N010", f["b2"].name)]["assignments_total"] == 0


async def test_b_students_per_trainer_lists_each_trainers_students(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    body = await _read(client, "students-by-trainer", course_id=str(f["program"].id))
    assert [c["key"] for c in body["columns"]] == ["trainer", "batch", "course", "student_code", "student", "status", "progress"]
    pairs = [(r["trainer"], r["student"]) for r in body["items"]]
    assert sorted(pairs, key=lambda p: (p[0].lower(), p[1].lower())) == pairs
    assert {p for p in pairs if p[0] == f["t2"].full_name} == {(f["t2"].full_name, "Aaron N010")}
    only = await _read(client, "students-by-trainer", trainer_id=str(f["t1"].id), course_id=str(f["program"].id))
    assert {r["student"] for r in only["items"]} == {"Aaron N010", "Bella N010", "Chen N010"}


async def test_c_course_completion_counts_completed_seats_per_batch(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    rows = _by(await _read(client, "course-completion", course_id=str(f["program"].id)), "batch")
    b1, b2 = rows[f["b1"].name], rows[f["b2"].name]
    assert (b1["students"], b1["completed"], b1["completion_pct"], b1["certificates"]) == (2, 1, 50.0, 1)  # the dropped seat is not counted
    assert (b2["students"], b2["completed"], b2["completion_pct"], b2["certificates"]) == (1, 0, 0.0, 0)
    assert b1["batch_ends"] == "2020-02-01" and list(rows) == [f["b2"].name, f["b1"].name]  # latest end first
    narrowed = await _read(client, "course-completion", course_id=str(f["program"].id), completion_to="2020-02-15")
    assert [r["batch"] for r in narrowed["items"]] == [f["b1"].name]


async def test_d_completions_in_a_period_lists_completed_seats_whose_batch_ends_in_it(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    body = await _read(client, "completions", course_id=str(f["program"].id), completion_from="2020-01-15", completion_to="2020-02-10")
    assert [(r["student"], r["batch"], r["attendance_pct"], r["certificate"]) for r in body["items"]] == [("Aaron N010", f["b1"].name, 66.67, "issued")]
    later = await _read(client, "completions", course_id=str(f["program"].id), completion_from="2020-02-15")
    assert later["items"] == []  # B2's only seat is at 50 %; the dropped seat never completes


async def test_e_attendance_counts_each_status_per_seat(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    body = await _read(client, "attendance", batch_id=str(f["b1"].id))
    assert [c["key"] for c in body["columns"]] == [
        "student_code", "student", "course", "batch", "trainer", "marked", "present", "absent", "late", "excused", "attendance_pct",
    ]
    rows = _by(body, "student")
    assert {k: rows["Aaron N010"][k] for k in ("marked", "present", "absent", "late", "excused", "attendance_pct")} == {
        "marked": 6, "present": 2, "absent": 2, "late": 1, "excused": 1, "attendance_pct": 66.67,
    }
    assert (rows["Bella N010"]["marked"], rows["Bella N010"]["attendance_pct"]) == (0, None)
    first_day = _by(await _read(client, "attendance", batch_id=str(f["b1"].id), **{"from": "2020-01-06", "to": "2020-01-07"}), "student")
    assert (first_day["Aaron N010"]["marked"], first_day["Aaron N010"]["attendance_pct"]) == (2, 100.0)


async def test_e_attendance_matches_certificate_eligibility(client, db_session):
    """AC2: the same seat's attendance % from the report and from the eligibility endpoint."""
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    eligibility = (await client.get(f"/api/v1/workflows/it/certificates/eligibility/{f['s1'].id}")).json()
    row = _by(await _read(client, "attendance", batch_id=str(f["b1"].id)), "student")["Aaron N010"]
    assert eligibility["attendance_percent"] == row["attendance_pct"] == 66.67
    assert (eligibility["assignments_graded"], eligibility["assignments_total"], eligibility["assessments_passed"], eligibility["assessments_total"]) == (1, 2, 1, 2)
    empty = (await client.get(f"/api/v1/workflows/it/certificates/eligibility/{f['s2'].id}")).json()
    assert empty["attendance_percent"] == 0  # unchanged: eligibility still reads an empty register as 0


async def test_f_trainer_report_summarises_each_trainer(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    rows = _by(await _read(client, "trainers", course_id=str(f["program"].id)), "trainer")
    t1, t2 = rows[f["t1"].full_name], rows[f["t2"].full_name]
    assert (t1["batches"], t1["students"], t1["attendance_pct"], t1["average_progress"], t1["pending_grading"]) == (1, 2, 66.67, 70.0, 1)
    assert (t2["batches"], t2["students"], t2["attendance_pct"], t2["average_progress"], t2["pending_grading"]) == (1, 1, None, 50.0, 0)


async def test_g_course_wise_report_with_revenue_for_super_admin(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "super_admin")
    row = (await _read(client, "courses", course_id=str(f["program"].id)))["items"]
    assert row == [{"course": f["program"].title, "batches": 2, "enrolments": 4, "active": 2, "completed": 1, "dropped": 1, "revenue": 1500.0}]
    none_enrolled = (await _read(client, "courses", course_id=str(f["program"].id), **{"from": "2021-01-01"}))["items"]
    assert none_enrolled == [{"course": f["program"].title, "batches": 2, "enrolments": 0, "active": 0, "completed": 0, "dropped": 0, "revenue": None}]


# --- access (AC3, AC4) -------------------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("role", ["it_admin", "trainer"])
async def test_g_revenue_is_left_out_without_the_financial_grant(client, db_session, role):
    f = await _fixture(db_session)
    user = f["t1"] if role == "trainer" else await make_user(db_session, role, "it")
    await login(client, user)
    body = await _read(client, "courses", course_id=str(f["program"].id))
    assert "revenue" not in [c["key"] for c in body["columns"]] and all("revenue" not in r for r in body["items"])
    file = await client.get(f"{API}/courses.xlsx", params={"course_id": str(f["program"].id)})
    assert file.status_code == 200
    header = [c.value for c in next(openpyxl.load_workbook(io.BytesIO(file.content)).active.iter_rows(max_row=1))]
    assert not any("Revenue" in str(h) for h in header)


@pytest.mark.parametrize("role", ["hr_team", "placement_team"])
async def test_recruiter_and_hr_cannot_read_the_pack(client, db_session, role):
    await _as(client, db_session, role)
    for report_id in PACK:
        assert (await client.get(f"{API}/{report_id}")).status_code == 403


async def test_a_trainer_sees_only_their_own_batches_in_every_report(client, db_session):
    f = await _fixture(db_session)
    await login(client, f["t1"])
    own = {"student-courses": 3, "students-by-trainer": 3, "course-completion": 1, "completions": 1, "attendance": 3, "trainers": 1, "courses": 1}
    for report_id, count in own.items():
        body = await _read(client, report_id, course_id=str(f["program"].id))
        assert body["total"] == count, report_id
        assert f["b2"].name not in {r.get("batch") for r in body["items"]} and f["t2"].full_name not in {r.get("trainer") for r in body["items"]}
    assert (await _read(client, "courses", course_id=str(f["program"].id)))["items"][0]["enrolments"] == 3
    for report_id in ("student-courses", "attendance", "trainers"):
        assert (await _read(client, report_id, trainer_id=str(f["t2"].id)))["items"] == []  # another trainer: empty, not 403


# --- downloads (AC1) ---------------------------------------------------------------------------------------------------------------


async def test_every_report_downloads_the_rows_of_its_json(client, db_session):
    f = await _fixture(db_session)
    await _as(client, db_session, "it_admin")
    for report_id in PACK:
        params = {"course_id": str(f["program"].id)}
        body = await _read(client, report_id, **params)
        file = await client.get(f"{API}/{report_id}.xlsx", params=params)
        assert file.status_code == 200, (report_id, file.text)
        sheet = list(openpyxl.load_workbook(io.BytesIO(file.content)).active.iter_rows(values_only=True))
        assert list(sheet[0]) == [c["label"] for c in body["columns"]], report_id
        assert len(sheet) - 1 == body["total"], report_id


async def test_an_undeclared_filter_is_refused(client, db_session):
    await _as(client, db_session, "it_admin")
    response = await client.get(f"{API}/trainers", params={"student": "x"})
    assert response.status_code == 422 and "Student" in response.json()["detail"]
