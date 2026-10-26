"""nehn-004 (DEC-SCOPE-280 X3-X5): every CSV report export has an `.xlsx` twin -- the same rows (typed), the same refusals, the same
audit row (AC1, AC2, AC5) -- for the telecaller reports and performance, the recruiter performance, the agency reports and commissions,
and the partnership reports."""

import csv
import io
from datetime import datetime

import openpyxl
import pytest
from sqlalchemy import select

from app.models import AuditLog
from app.services.tabular import XLSX_TYPE
from tests.agn001_helpers import client_for
from tests.agn014_helpers import mk_commission
from tests.agn020_helpers import REPORTS as AGENT_REPORTS
from tests.agn020_helpers import reports_world
from tests.bdm017_helpers import as_user
from tests.rec001_helpers import as_role as rec_as_role
from tests.rec001_helpers import login as rec_login
from tests.rec001_helpers import make_pm as rec_make_pm
from tests.tel004_helpers import make_telecaller, make_tl_manager
from tests.upc003_helpers import as_role as upc_as_role
from tests.upc003_helpers import login as upc_login
from tests.upc003_helpers import make_head

pytestmark = pytest.mark.asyncio
API = "/api/v1"
COMMISSIONS = f"{API}/workflows/overseas/agent/commissions/report"


def _csv_rows(response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def _same(value, text: str) -> bool:
    """An xlsx cell equals the CSV cell for the same payload value: a number by value (the CSV writes `0.0`, `12.50`), a day as ISO."""
    if value is None:
        return text == ""
    if isinstance(value, datetime):
        return text == (value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ", "minutes"))
    if isinstance(value, int | float):
        return float(text) == value
    return str(value) == text


def _xlsx_rows(response) -> list[list]:
    return [list(row) for row in openpyxl.load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True)]


def _same_table(xlsx: list[list], csv_rows: list[list[str]]) -> bool:
    return len(xlsx) == len(csv_rows) and all(len(a) == len(b) and all(map(_same, a, b)) for a, b in zip(xlsx, csv_rows, strict=True))


async def _audits(db, user, action: str) -> list[AuditLog]:
    stmt = select(AuditLog).where(AuditLog.user_id == user.id, AuditLog.action == action).order_by(AuditLog.created_at)
    return list((await db.scalars(stmt.execution_options(populate_existing=True))).all())


async def _twins(client, path: str, params: dict | None = None):
    csv_response = await client.get(f"{path}.csv", params=params or {})
    xlsx_response = await client.get(f"{path}.xlsx", params=params or {})
    assert csv_response.status_code == 200, csv_response.text
    assert xlsx_response.status_code == 200, xlsx_response.text
    assert xlsx_response.headers["content-type"] == XLSX_TYPE
    assert xlsx_response.headers["cache-control"] == "private, no-store"
    disposition = csv_response.headers["content-disposition"].replace(".csv", ".xlsx")
    assert xlsx_response.headers["content-disposition"] == disposition
    assert _same_table(_xlsx_rows(xlsx_response), _csv_rows(csv_response)), (_xlsx_rows(xlsx_response), _csv_rows(csv_response))  # AC1
    return csv_response, xlsx_response


def _assert_audited(rows: list[AuditLog]) -> None:
    """AC5: one audit row per download, CSV first then xlsx, the format recorded."""
    assert [row.metadata_json["format"] for row in rows] == ["csv", "xlsx"]
    assert {k: v for k, v in rows[0].metadata_json.items() if k != "format"} == {k: v for k, v in rows[1].metadata_json.items() if k != "format"}


# --- telecaller (tel-024, tel-023) --------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("kind", ["source", "product", "telecaller", "handover", "campaign"])
async def test_telecaller_reports(client, db_session, kind):
    manager = await make_tl_manager(db_session)
    await make_telecaller(db_session, manager)
    await as_user(client, manager)
    await _twins(client, f"{API}/telecaller/reports/{kind}")
    _assert_audited(await _audits(db_session, manager, "telecaller_report.export"))


async def test_telecaller_reports_refusals_match(client, db_session):
    manager = await make_tl_manager(db_session)
    await as_user(client, await make_telecaller(db_session, manager))
    assert (await client.get(f"{API}/telecaller/reports/source.xlsx")).status_code == 403
    await as_user(client, manager)
    assert (await client.get(f"{API}/telecaller/reports/nope.xlsx")).status_code == 404
    response = await client.get(f"{API}/telecaller/reports/source.xlsx", params={"date_from": "nope"})
    assert response.status_code == 422 and response.json() == (await client.get(f"{API}/telecaller/reports/source.csv", params={"date_from": "nope"})).json()


async def test_telecaller_performance(client, db_session):
    manager = await make_tl_manager(db_session)
    tel = await make_telecaller(db_session, manager)
    await as_user(client, manager)
    await _twins(client, f"{API}/telecaller/manager/performance")
    _assert_audited(await _audits(db_session, manager, "telecaller_performance.export"))
    await as_user(client, tel)
    assert (await client.get(f"{API}/telecaller/manager/performance.xlsx")).status_code == 403


# --- recruiter (rec-033) ------------------------------------------------------------------------------------------------------------
async def test_recruiter_performance(client, db_session):
    manager = await rec_make_pm(db_session)
    await rec_login(client, manager)
    await _twins(client, f"{API}/recruiter/performance")
    _assert_audited(await _audits(db_session, manager, "recruiter_performance.export"))
    await rec_as_role(client, db_session, "bdm", "it")
    assert (await client.get(f"{API}/recruiter/performance.xlsx")).status_code == 403


# --- partnership (upc-031) ----------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("kind", ["pipeline", "expected", "performance", "agreements", "targets"])
async def test_partnership_reports(client, db_session, kind):
    head = await make_head(db_session)
    await upc_login(client, head)
    await _twins(client, f"{API}/partnership/reports/{kind}")
    _assert_audited(await _audits(db_session, head, "partnership_report.export"))


async def test_partnership_refusals_match(client, db_session):
    assert (await client.get(f"{API}/partnership/reports/pipeline.xlsx")).status_code == 401
    await upc_as_role(client, db_session, "counselor", "overseas")
    assert (await client.get(f"{API}/partnership/reports/nonsense.xlsx")).status_code == 403
    await upc_login(client, await make_head(db_session))
    assert (await client.get(f"{API}/partnership/reports/nonsense.xlsx")).status_code == 404


# --- agency (agn-020, agn-014) ------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("kind", ["students", "applications", "enrollments", "universities", "countries", "intakes", "staff"])
async def test_agency_reports(db_session, kind):
    world = await reports_world(db_session)
    async with client_for(world["master"].email) as c:
        csv_response, _ = await _twins(c, f"{AGENT_REPORTS}/{kind}")
    _assert_audited(await _audits(db_session, world["master"], "agent_report.export"))
    assert len(_csv_rows(csv_response)) > 1


async def test_agency_dates_are_typed_and_names_stay_literal(db_session):
    world = await reports_world(db_session)
    async with client_for(world["master"].email) as c:
        response = await c.get(f"{AGENT_REPORTS}/applications.xlsx")
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    typed = [cell for row in sheet.iter_rows(min_row=2) for cell in row if cell.is_date]
    assert typed, "a day column comes out as a date cell (AC2)"
    assert all(cell.data_type == "s" for row in sheet.iter_rows(min_row=2) for cell in row[:1])  # names are text, never formulas


async def test_agency_staff_report_stays_master_only_and_the_cap_holds(db_session, monkeypatch):
    world = await reports_world(db_session)
    async with client_for(world["s1"]["user"].email) as c:
        assert (await c.get(f"{AGENT_REPORTS}/staff.xlsx")).status_code == 403
    from app.services import agent_reports

    monkeypatch.setattr(agent_reports, "CSV_ROW_CAP", 1)
    async with client_for(world["master"].email) as c:
        response = await c.get(f"{AGENT_REPORTS}/students.xlsx")
    assert response.status_code == 422 and "more than 1 rows" in response.json()["detail"]
    assert await _audits(db_session, world["master"], "agent_report.export") == []  # a refused export writes nothing


async def test_agency_commissions(db_session):
    world = await reports_world(db_session)
    await mk_commission(db_session, world, amount=1234.5, student_name="=cmd|' /C calc'!A0")
    async with client_for(world["master"].email) as c:
        _, response = await _twins(c, COMMISSIONS)
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    row = next(r for r in sheet.iter_rows(min_row=2) if r[0].value.startswith("'=cmd"))
    assert row[0].data_type == "s"  # AC3: literal text
    assert row[5].value == 1234.5 and row[5].data_type == "n"  # AC2: the amount is a number
    assert row[7].is_date  # AC2: Created is a date
    async with client_for(world["s1"]["user"].email) as c:
        assert (await c.get(f"{COMMISSIONS}.xlsx")).status_code == (await c.get(f"{COMMISSIONS}.csv")).status_code
