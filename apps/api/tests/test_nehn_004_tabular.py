"""nehn-004 (DEC-SCOPE-280): the shared tabular helper -- the xlsx writer's typed dates, header style, auto-filter and widths (X1), and
`read_tabular`, the hardened `.xlsx`/`.csv` reader every import can share (X2)."""

import io
import zipfile
from datetime import UTC, date, datetime
from decimal import Decimal
from xml.etree import ElementTree

import openpyxl
import pytest

from app.services.tabular import TabularError, read_tabular, to_xlsx

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
COLUMNS = [{"key": "name", "label": "Name"}, {"key": "day", "label": "Day"}, {"key": "phone", "label": "Phone"}, {"key": "amount", "label": "Amount"}]
PAYLOAD = {
    "columns": COLUMNS,
    "items": [
        {"name": "தமிழ் हिन्दी", "day": "2026-10-26", "phone": "09876543210", "amount": Decimal("12.50")},
        {"name": '=HYPERLINK("http://x")', "day": date(2026, 1, 2), "phone": "+91 1", "amount": 3},
        {"name": "Late", "day": datetime(2026, 10, 26, 18, 45, tzinfo=UTC), "phone": None, "amount": 1.5},
    ],
    "totals": {"name": "Total", "day": None, "phone": None, "amount": Decimal("17.00")},
}


def _book(body: bytes):
    return openpyxl.load_workbook(io.BytesIO(body))


def _part(body: bytes, name: str):
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        return ElementTree.fromstring(archive.read(name))


# --- writer (X1) -------------------------------------------------------------------------------------------------------------------
def test_dates_and_numbers_are_typed_cells_and_text_stays_text():
    sheet = _book(to_xlsx(PAYLOAD, "Report")).active
    assert sheet["B2"].value == datetime(2026, 10, 26) and sheet["B2"].is_date  # an ISO day string in the payload
    assert sheet["B3"].value == datetime(2026, 1, 2) and sheet["B3"].is_date
    assert sheet["B4"].value == datetime(2026, 10, 27, 0, 15) and sheet["B4"].number_format == "yyyy-mm-dd hh:mm"  # UTC -> IST
    assert sheet["D2"].value == 12.5 and sheet["D3"].value == 3 and sheet["D5"].value == 17
    assert sheet["C2"].value == "09876543210"  # leading zero kept: text, never a number
    assert sheet["A2"].value == "தமிழ் हिन्दी"


def test_a_formula_looking_value_is_literal_text():
    sheet = _book(to_xlsx(PAYLOAD, "Report")).active
    assert sheet["A3"].data_type == "s" and sheet["A3"].value == '\'=HYPERLINK("http://x")'  # AC3


def test_the_header_is_bold_frozen_filtered_and_columns_have_widths():
    body = to_xlsx(PAYLOAD, "Report")
    sheet = _book(body).active
    assert all(cell.font.b for cell in sheet[1]) and not sheet["A2"].font.b
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:D4"  # header + items; the Total row stays out of sorting
    assert 8 <= sheet.column_dimensions["A"].width <= 60 and sheet.column_dimensions["A"].width >= len('\'=HYPERLINK("http://x")')


def test_the_package_declares_its_styles_part():
    body = to_xlsx(PAYLOAD, "Report")
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        assert "xl/styles.xml" in archive.namelist()
    types = _part(body, "[Content_Types].xml")
    assert any(o.get("PartName") == "/xl/styles.xml" for o in types)


def test_an_empty_report_is_a_header_only_sheet_with_a_filter():
    sheet = _book(to_xlsx({"columns": COLUMNS, "items": [], "totals": None}, "Empty")).active
    assert [c.value for c in sheet[1]] == ["Name", "Day", "Phone", "Amount"] and sheet.max_row == 1
    assert sheet.auto_filter.ref == "A1:D1"


def test_a_non_date_string_that_looks_close_stays_text():
    payload = {"columns": COLUMNS[:1], "items": [{"name": "2026-13-40"}, {"name": "2026-10-26 extra"}], "totals": None}
    sheet = _book(to_xlsx(payload, "R")).active
    assert sheet["A2"].value == "2026-13-40" and sheet["A3"].value == "2026-10-26 extra"


# --- reader (X2) -------------------------------------------------------------------------------------------------------------------
def _xlsx(rows: list[list]) -> bytes:
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_reads_an_xlsx_into_rows_by_header_skipping_blank_rows():
    raw = _xlsx([[" Question ", "Answer", "Marks"], ["தமிழ்?", "ஆம்", 2], [None, None, None], ["Q2", "0123", 1.0]])
    rows = read_tabular("Questions.XLSX", raw, required=("question", "answer"))
    assert rows == [(2, {"question": "தமிழ்?", "answer": "ஆம்", "marks": "2"}), (4, {"question": "Q2", "answer": "0123", "marks": "1"})]


def test_xlsx_dates_come_back_as_iso_text():
    raw = _xlsx([["name", "day", "at"], ["a", date(2026, 10, 26), datetime(2026, 10, 26, 9, 30)], ["b", datetime(2026, 1, 1), True]])
    rows = read_tabular("d.xlsx", raw, required=("name",))
    assert rows[0][1] == {"name": "a", "day": "2026-10-26", "at": "2026-10-26 09:30"}
    assert rows[1][1] == {"name": "b", "day": "2026-01-01", "at": "TRUE"}


def test_reads_a_csv_with_a_bom():
    raw = "﻿Question,Answer\r\nQ1, A1 \r\n,\r\n".encode()
    assert read_tabular("q.csv", raw, required=("question",)) == [(2, {"question": "Q1", "answer": "A1"})]


@pytest.mark.parametrize("name", ["q.xlsx", "q.csv"])
def test_every_missing_required_column_gets_its_own_message(name):
    raw = _xlsx([["Marks"], [1]]) if name.endswith("xlsx") else b"Marks\r\n1\r\n"
    with pytest.raises(TabularError) as error:
        read_tabular(name, raw, required=("question", "answer"))
    assert error.value.status_code == 422
    assert error.value.messages == ["Missing required column: question", "Missing required column: answer"]  # AC4


def test_unknown_and_repeated_columns_are_reported_when_the_columns_are_known():
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", _xlsx([["question", "Question", "colour"], ["a", "b", "c"]]), required=("question",), known=("question",))
    assert error.value.messages == ["Duplicate column: question", "Unknown column: colour"]


@pytest.mark.parametrize("name", ["q.xlsm", "q.xls", "q.txt", "q"])
def test_only_xlsx_and_csv_are_accepted(name):
    with pytest.raises(TabularError) as error:
        read_tabular(name, _xlsx([["question"], ["a"]]), required=("question",))
    assert error.value.status_code == 415


def test_a_macro_workbook_renamed_to_xlsx_is_refused():
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_xlsx([["question"], ["a"]]))) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"\x00macro")
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", buffer.getvalue(), required=("question",))
    assert error.value.status_code == 415


def test_a_dtd_in_any_part_is_refused():
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_xlsx([["question"], ["a"]]))) as source, zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/workbook.xml":
                data = b'<!DOCTYPE x [<!ENTITY a "aaaa">]>' + data.split(b"?>", 1)[-1]  # openpyxl may omit the declaration
            target.writestr(item, data)
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", buffer.getvalue(), required=("question",))
    assert error.value.status_code == 422


def test_a_zip_bomb_is_refused_before_it_is_inflated():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * (25 * 1024 * 1024))
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", buffer.getvalue(), required=("question",))
    assert error.value.status_code == 422


@pytest.mark.parametrize("raw", [b"not a zip", b"PK\x03\x04broken"])
def test_a_corrupt_file_is_refused(raw):
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", raw, required=("question",))
    assert error.value.status_code == 422 and error.value.messages == ["The file is not a valid Excel (.xlsx) file"]


def test_size_and_row_caps():
    raw = _xlsx([["question"], *[[f"q{i}"] for i in range(4)]])
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", raw, required=("question",), max_bytes=100)
    assert error.value.status_code == 413
    with pytest.raises(TabularError) as error:
        read_tabular("q.xlsx", raw, required=("question",), max_rows=3)
    assert error.value.status_code == 422 and error.value.messages == ["The file has more than 3 rows"]
    assert len(read_tabular("q.xlsx", raw, required=("question",), max_rows=4)) == 4


def test_a_formula_cell_is_never_evaluated():
    book = openpyxl.Workbook()
    book.active.append(["question"])
    book.active.append(['=HYPERLINK("http://evil","x")'])
    buffer = io.BytesIO()
    book.save(buffer)
    # data_only: the formula is not run; a file never saved by Excel has no cached value, so the cell is empty
    assert read_tabular("q.xlsx", buffer.getvalue(), required=("question",)) == []
